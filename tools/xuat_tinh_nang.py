"""Xuất 120 tính năng IACC Cloud ra src/app/features.json cho web.

Đọc qua build_present.load_features() của bộ present để dùng chung GOI_FIX, EXTRA, KHO_IVT, KHONG_IVT.
Không sửa gì bên Present. Excel đổi thì chạy lại:  python tools/xuat_tinh_nang.py
"""
import contextlib
import io
import json
import sys
from pathlib import Path

import openpyxl

WEB = Path(__file__).resolve().parent.parent
PRESENT = WEB.parent / 'Present' / 'tools'
sys.path.insert(0, str(PRESENT))
with contextlib.redirect_stdout(io.StringIO()):
    import build_present as b
    feats = b.load_features()

# Tên nhóm con (Chứng từ, Sổ sách...) lấy từ dòng Excel cấp 2
ws = openpyxl.load_workbook(b.XLSX, data_only=True)['RoadMap']
names = {str(ws[f'A{r}'].value or '').strip(): b.norm(ws[f'B{r}'].value)
         for r in range(4, ws.max_row + 1) if ws[f'B{r}'].value}
GRP = {'Sổ sách báo cáo': 'Sổ sách, báo cáo', 'Sổ sách và báo cáo': 'Sổ sách, báo cáo'}
FIRST = {'Danh mục': 'Danh mục', 'Tiện ích bổ sung': 'Tiện ích'}

out = []
for f in feats:
    code = f['code'].split('-')[0]
    parts = code.split('.')
    top = b.TOP.get(parts[0])
    if f['mod'] in FIRST:
        grp = FIRST[f['mod']]
    elif len(parts) == 3 and top == f['mod']:
        grp = names[parts[0] + '.' + parts[1]]
    else:
        grp = 'Chứng từ'          # dòng chuyển phân hệ, như Công thức chế biến sang Kho hàng
    out.append(dict(c=f['code'], m=b.MODS.index(f['mod']), n=f['name'], grp=GRP.get(grp, grp),
                    g=''.join(f['goi']), gd=f['gd'], ivt=int(f['ivt'])))

dst = WEB / 'src' / 'app' / 'features.json'
dst.write_text(json.dumps(dict(mods=b.MODS, feats=out), ensure_ascii=False, indent=0), encoding='utf-8')
cnt = {g: sum(g in x['g'] for x in out) for g in 'FSMA'}
print(f'{len(out)} tính năng -> {dst}  theo gói: {cnt}')
